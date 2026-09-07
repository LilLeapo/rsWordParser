"""Read-only rendered-PDF, chart-cache and embedded-workbook audit for UI samples."""

import argparse
import hashlib
import importlib.util
import io
import itertools
import json
import posixpath
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from lxml import etree
from openpyxl.utils.cell import range_to_tuple


ROOT = Path('C:/word/real-word-round3-20260907')
SOURCE = Path('C:/word/round3-work-20260907/real-word-round3-inputs/edited')
NS = {
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
    'w14': 'http://schemas.microsoft.com/office/word/2010/wordml',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
}


def xp(node, expression):
    return node.xpath(expression, namespaces=NS)


def digest(data):
    return hashlib.sha256(data).hexdigest().upper()


def value(text):
    if text is None:
        return None
    try:
        numeric = float(text)
        return int(numeric) if numeric.is_integer() else numeric
    except (ValueError, TypeError):
        return text


def scalar_text(nodes):
    return ''.join(str(item) for item in nodes)


def resolve_target(part, target):
    if target.startswith('/'):
        return target.lstrip('/')
    return posixpath.normpath(posixpath.join(posixpath.dirname(part), target))


def inspect_drawing_ids(doc):
    tree = doc.getroottree()
    occurrences = []
    for node in xp(doc, '//wp:docPr'):
        branches = {}
        for ancestor in node.iterancestors():
            if ancestor.tag in ('{' + NS['mc'] + '}Choice', '{' + NS['mc'] + '}Fallback'):
                branches[tree.getpath(ancestor.getparent())] = tree.getpath(ancestor)
        occurrences.append({'id': node.get('id'), 'path': tree.getpath(node), 'mceBranches': branches})
    collisions = []
    for left, right in itertools.combinations(occurrences, 2):
        if left['id'] != right['id']:
            continue
        shared = left['mceBranches'].keys() & right['mceBranches'].keys()
        exclusive = any(left['mceBranches'][key] != right['mceBranches'][key] for key in shared)
        if not exclusive:
            collisions.append({'id': left['id'], 'paths': [left['path'], right['path']]})
    return occurrences, collisions


def inspect_package(path):
    raw = path.read_bytes()
    result = {'path': str(path), 'sha256': digest(raw), 'charts': []}
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        result['crcError'] = archive.testzip()
        names = archive.namelist()
        result['duplicateMembers'] = len(names) - len(set(names))
        parts = {name: etree.fromstring(archive.read(name)) for name in names if name.endswith(('.xml', '.rels'))}
        doc = parts['word/document.xml']
        ids = xp(doc, '//wp:docPr/@id')
        result['docPrIdsAllMceBranches'] = ids
        result['docPrIdsUnique'] = len(ids) == len(set(ids))
        result['docPrOccurrences'], result['docPrPotentialActiveCollisions'] = inspect_drawing_ids(doc)
        result['docPrIdsUniqueWithinAnyMceSelection'] = not result['docPrPotentialActiveCollisions']
        result['nativeInkPartHashes'] = {name: digest(archive.read(name)) for name in names if name.startswith('word/ink/')}
        result['nativeContentPartCount'] = len(xp(doc, '//w14:contentPart'))
        result['mediaHashes'] = {name: digest(archive.read(name)) for name in names if name.startswith('word/media/')}
        for part_name, chart in parts.items():
            if not part_name.startswith('word/charts/') or etree.QName(chart).localname != 'chartSpace':
                continue
            if etree.QName(chart).namespace != NS['c']:
                result['charts'].append({'part': part_name, 'namespace': etree.QName(chart).namespace,
                                         'sha256': digest(archive.read(part_name)),
                                         'unresolved': 'Non-classic chart namespace; preserve without claiming cache/workbook comparison.'})
                continue
            rel_name = posixpath.join(posixpath.dirname(part_name), '_rels', posixpath.basename(part_name) + '.rels')
            rels = {node.get('Id'): node for node in parts.get(rel_name, [])}
            workbook_ids = xp(chart, 'c:externalData/@r:id')
            workbook_path, workbook = None, None
            external = []
            for rid in workbook_ids:
                rel = rels.get(rid)
                if rel is None:
                    external.append({'id': rid, 'error': 'Missing chart relationship'})
                    continue
                target = resolve_target(part_name, rel.get('Target', ''))
                if rel.get('TargetMode') == 'External' or target not in names:
                    external.append({'id': rid, 'target': target, 'unresolved': 'External or absent workbook; not accessed'})
                    continue
                workbook_path = target
                workbook_raw = archive.read(target)
                item = {'id': rid, 'part': target, 'sha256': digest(workbook_raw)}
                try:
                    workbook = openpyxl.load_workbook(io.BytesIO(workbook_raw), read_only=True, data_only=True)
                    item['sheets'] = {sheet.title: [[cell.value for cell in row] for row in sheet.iter_rows()]
                                      for sheet in workbook.worksheets}
                except Exception as error:
                    item['error'] = str(error)
                external.append(item)
            series = []
            for index, ser in enumerate(xp(chart, '//c:ser'), 1):
                entry = {'index': index, 'channels': []}
                for child in ser:
                    name = etree.QName(child).localname
                    if name not in ('tx', 'cat', 'val', 'xVal', 'yVal', 'bubbleSize'):
                        continue
                    formula = scalar_text(xp(child, './/c:f/text()')) or None
                    cache = [{'index': int(point.get('idx')), 'value': value(scalar_text(xp(point, 'c:v/text()')))}
                             for point in xp(child, './/c:numCache/c:pt | .//c:strCache/c:pt | .//c:numLit/c:pt | .//c:strLit/c:pt')]
                    cache.sort(key=lambda item: item['index'])
                    direct = xp(child, 'c:v/text()')
                    data = {'channel': name, 'formula': formula, 'cachedPoints': cache,
                            'cachedValues': [item['value'] for item in cache] if cache else [value(item) for item in direct]}
                    if formula and workbook:
                        try:
                            sheet_name, bounds = range_to_tuple(formula)
                            min_col, min_row, max_col, max_row = bounds
                            values = [cell.value for row in workbook[sheet_name].iter_rows(min_row=min_row, max_row=max_row,
                                      min_col=min_col, max_col=max_col) for cell in row]
                            data['embeddedWorkbookValues'] = values
                            data['cacheEqualsEmbeddedWorkbook'] = data['cachedValues'] == values
                        except Exception as error:
                            data['referenceReadError'] = str(error)
                    entry['channels'].append(data)
                series.append(entry)
            result['charts'].append({'part': part_name, 'sha256': digest(archive.read(part_name)), 'title': scalar_text(xp(chart, 'c:chart/c:title//a:t/text()')),
                                     'workbooks': external, 'series': series})
            if workbook:
                workbook.close()
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--file', action='append', help='UI input filename; repeat to select. Default: every completed UI row.')
    parser.add_argument('--no-pdf', action='store_true')
    args = parser.parse_args()
    selected = set(args.file or [])
    ui = json.loads((ROOT / 'ui-edited3.json').read_text(encoding='utf-8-sig'))
    ui = [item for item in ui if not selected or item['file'] in selected]
    annotations_path = ROOT / '_scripts/edited3-pdf-annotations.json'
    annotations = json.loads(annotations_path.read_text(encoding='utf-8-sig')) if annotations_path.exists() else {}
    spec = importlib.util.spec_from_file_location('revfix', ROOT / '_scripts/inspect-revfix.py')
    revfix = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(revfix)
    stamp = datetime.now(timezone.utc).isoformat()
    pdf_report = {'generatedUtc': stamp, 'method': 'Independent rendered Word-exported PDF inspection. PDF is exported after UI inspection, workbook activation where applicable and native Word resave. It is not proof of absence of a recovery dialog or of pre-activation chart values.', 'rows': []}
    chart_report = {'generatedUtc': stamp, 'method': 'Read-only ZIP/lxml chart-cache inspection and openpyxl reading of embedded workbook cells. Source and resaved DOCX are never opened in Word by this inspector or modified. Cache/workbook comparison does not simulate Word activation.', 'rows': []}
    for item in ui:
        filename = item['file']
        stem = Path(filename).stem
        source = inspect_package(SOURCE / filename)
        resaved = inspect_package(Path(item['resaved']))
        assert resaved['sha256'] == item['resaved_sha256'], 'Resaved DOCX hash changed: ' + filename
        evidence = {'file': filename, 'source': source, 'resaved': resaved,
                    'sameNativeInkBytes': sorted(source['nativeInkPartHashes'].values()) == sorted(resaved['nativeInkPartHashes'].values()),
                    'originalInkCount': len(source['nativeInkPartHashes']),
                    'observedByWordUi': {'readout': item['initialReadout'], 'activationReadout': item['activationReadout']}}
        chart_report['rows'].append(evidence)
        if not args.no_pdf:
            pdf = revfix.inspect_pdf(Path(item['pdf']), ROOT / 'screenshots/edited3-pdf', stem, annotations.get(filename))
            pdf_report['rows'].append({'file': filename, 'sourceSha256': source['sha256'], 'resavedSha256': resaved['sha256'], 'pdf': pdf})
    pdf_report['summary'] = {'pdfs': len(pdf_report['rows']), 'pages': sum(row['pdf']['pageCount'] for row in pdf_report['rows']),
                             'visuallyReviewed': sum(row['pdf']['visualReviewed'] for row in pdf_report['rows'])}
    chart_report['summary'] = {'files': len(chart_report['rows']),
                               'filesWithCharts': sum(bool(row['source']['charts']) for row in chart_report['rows']),
                               'sourceAndResavedPackagesInspected': len(chart_report['rows']) * 2,
                               'zipIntegrityIssues': [{'file': row['file'], 'stage': stage, 'crcError': row[stage]['crcError'], 'duplicateMembers': row[stage]['duplicateMembers']}
                                                      for row in chart_report['rows'] for stage in ('source', 'resaved')
                                                      if row[stage]['crcError'] or row[stage]['duplicateMembers']],
                               'potentialActiveDrawingIdCollisions': [{'file': row['file'], 'stage': stage, 'collisions': row[stage]['docPrPotentialActiveCollisions']}
                                                                    for row in chart_report['rows'] for stage in ('source', 'resaved')
                                                                    if row[stage]['docPrPotentialActiveCollisions']],
                               'rawDuplicateIdsOnlyInExclusiveMceBranches': [{'file': row['file'], 'stage': stage, 'ids': row[stage]['docPrIdsAllMceBranches']}
                                                                          for row in chart_report['rows'] for stage in ('source', 'resaved')
                                                                          if not row[stage]['docPrIdsUnique'] and row[stage]['docPrIdsUniqueWithinAnyMceSelection']],
                               'nativeInkChangedFiles': [row['file'] for row in chart_report['rows'] if not row['sameNativeInkBytes']],
                               'sourceCacheWorkbookMismatches': [{'file': row['file'], 'part': chart['part'], 'series': ser['index'], 'channel': channel['channel'],
                                                                 'cache': channel['cachedValues'], 'workbook': channel['embeddedWorkbookValues']}
                                                                for row in chart_report['rows'] for chart in row['source']['charts'] for ser in chart.get('series', [])
                                                                for channel in ser['channels'] if channel.get('cacheEqualsEmbeddedWorkbook') is False],
                               'resavedCacheWorkbookMismatches': [{'file': row['file'], 'part': chart['part'], 'series': ser['index'], 'channel': channel['channel'],
                                                                  'cache': channel['cachedValues'], 'workbook': channel['embeddedWorkbookValues']}
                                                                 for row in chart_report['rows'] for chart in row['resaved']['charts'] for ser in chart.get('series', [])
                                                                 for channel in ser['channels'] if channel.get('cacheEqualsEmbeddedWorkbook') is False]}
    (ROOT / '_scripts/edited3-pdf-visual.json').write_text(json.dumps(pdf_report, ensure_ascii=True, indent=2), encoding='utf-8')
    (ROOT / '_scripts/edited3-chart-audit.json').write_text(json.dumps(chart_report, ensure_ascii=True, indent=2), encoding='utf-8')
    print(json.dumps({'pdfSummary': pdf_report['summary'],
                      'chartSummary': {key: (len(value) if isinstance(value, list) else value)
                                       for key, value in chart_report['summary'].items()}}, ensure_ascii=True, indent=2))


if __name__ == '__main__':
    main()
